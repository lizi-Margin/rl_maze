"""将 REPORT.md 转换为排版好的 .docx 作业报告。

流程：
1. pandoc 把 REPORT.md 转成 docx（LaTeX 公式转 Word 原生公式，
   相对路径 results/fig*.png 会被嵌入文档）；
2. python-docx 后处理：设置中文字体（正文宋体、标题微软雅黑）、
   正文 12pt（小四）、图片等比缩放到页宽内并居中。

用法：python make_docx.py   （输出 RL迷宫作业报告.docx）
"""
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
SRC = ROOT / 'REPORT.md'
OUT = ROOT / 'RL迷宫作业报告.docx'


def convert_with_pandoc():
    cmd = ['pandoc', str(SRC), '-o', str(OUT),
           '--from', 'markdown', '--to', 'docx']
    res = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True)
    if res.returncode != 0:
        sys.exit('pandoc 转换失败:\n' + res.stderr)
    if res.stderr.strip():
        print('pandoc 警告:\n' + res.stderr.strip())


def set_style_font(doc, style_name, east_asia, latin=None, size=None):
    """给段落样式设置中西文字体（w:eastAsia 属性控制中文字体）。"""
    try:
        style = doc.styles[style_name]
    except KeyError:
        return
    if latin:
        style.font.name = latin          # w:ascii / w:hAnsi
    if size:
        style.font.size = size
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn('w:eastAsia'), east_asia)


def postprocess():
    doc = Document(str(OUT))

    # ---- 字体 ----
    set_style_font(doc, 'Normal', '宋体', 'Times New Roman', Pt(12))
    set_style_font(doc, 'Title', '微软雅黑')
    for i in range(1, 5):
        set_style_font(doc, f'Heading {i}', '微软雅黑')
    set_style_font(doc, 'Image Caption', '宋体', 'Times New Roman', Pt(10))
    set_style_font(doc, 'Table Caption', '宋体', 'Times New Roman', Pt(10))

    # ---- 图片：等比缩放到页宽内 ----
    max_w = Inches(6.0)
    for shape in doc.inline_shapes:
        if shape.width > max_w:
            ratio = max_w / shape.width
            shape.height = round(shape.height * ratio)
            shape.width = round(max_w)

    # ---- 仅含图片的段落居中 ----
    for para in doc.paragraphs:
        has_drawing = any(
            run._element.findall(qn('w:drawing')) for run in para.runs)
        if has_drawing:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(OUT))


def verify():
    doc = Document(str(OUT))
    n_para = len(doc.paragraphs)
    n_table = len(doc.tables)
    n_img = len(doc.inline_shapes)
    size_kb = OUT.stat().st_size / 1024
    print(f'生成 {OUT.name}: {n_para} 段落, {n_table} 表格, {n_img} 张图, '
          f'{size_kb:.0f} KB')
    if n_img < 8:
        sys.exit('警告: 嵌入图片少于 8 张，请检查 results/ 路径')


if __name__ == '__main__':
    convert_with_pandoc()
    postprocess()
    verify()
    print('完成:', OUT)
